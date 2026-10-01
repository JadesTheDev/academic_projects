# SDC205L - Mr. Zakaria
# SDC205 Project
# Week 5 - Graphing Dynamically Generated Data

from datetime import datetime
import os
import csv

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference


# convertData requires value (float) and data_type (string).
# It converts the original value and returns the converted numerical value.
def convertData(value, data_type):
    if data_type == "1":
        # Fahrenheit to Celsius
        return (value - 32) * 5 / 9

    elif data_type == "2":
        # Pounds to kilograms
        return value / 2.205

    elif data_type == "3":
        # Inches to centimeters
        return value * 2.54


# getDataTypeName requires data_type (string).
# It returns the name of the selected data type.
def getDataTypeName(data_type):
    if data_type == "1":
        return "Temperature"

    elif data_type == "2":
        return "Weight"

    elif data_type == "3":
        return "Rain Amount"


# insertData requires file_path (string) and data (list).
# It writes the data to the CSV file and returns True or False.
def insertData(file_path, data):
    try:
        file_exists = os.path.exists(file_path)

        with open(file_path, "a", newline="") as file:
            writer = csv.writer(file)

            if not file_exists:
                writer.writerow([
                    "Date",
                    "Location",
                    "Data Type",
                    "Description",
                    "Original Value",
                    "Converted Value"
                ])

            writer.writerow(data)

        return True

    except OSError as error:
        print("Error writing to file:", error)
        return False


# viewData requires file_path (string).
# It displays the contents of the CSV file and returns no value.
def viewData(file_path):
    try:
        with open(file_path, "r") as file:
            print("\nThe file", os.path.abspath(file_path))

            for line in file:
                print(line.strip())

    except OSError as error:
        print("Error reading file:", error)


# getInput requires no arguments.
# It collects user data, saves it to ZooData.csv, and returns no value.
def getInput():
    try:
        number_of_entries = int(
            input("How many entries would you like to add? ")
        )

        for entry_number in range(number_of_entries):

            print(
                "\nEntry",
                entry_number + 1,
                "of",
                number_of_entries
            )

            date = input("Enter a date: ")
            location = input("Enter a location: ")
            description = input("Enter a description: ")

            print("\n1 Temperature")
            print("2 Weight")
            print("3 Rain Amount")

            data_type = input("Choose the type of data: ")

            if data_type not in ("1", "2", "3"):
                print("Error: Invalid data type selected.")
                continue

            data_type_name = getDataTypeName(data_type)

            if data_type == "1":
                value = float(
                    input("Enter temperature in Fahrenheit: ")
                )

            elif data_type == "2":
                value = float(
                    input("Enter weight in pounds: ")
                )

            else:
                value = float(
                    input("Enter rain amount in inches: ")
                )

            converted_value = convertData(
                value,
                data_type
            )

            data = [
                date,
                location,
                data_type_name,
                description,
                value,
                converted_value
            ]

            if insertData("ZooData.csv", data):
                print(
                    "The following data was saved at",
                    datetime.now(),
                    ":",
                    ",".join(str(item) for item in data)
                )

    except ValueError:
        print("Error: Please enter a valid number.")


# createChart requires file_path (string) and chart_type (string).
# It reads CSV data, creates final.xlsx with a chart, and returns no value.
def createChart(file_path, chart_type):
    try:
        print("\nChoose the type of data to graph")
        print("1 Temperature")
        print("2 Weight")
        print("3 Rain Amount")

        type_choice = input("Choose an option: ")

        if type_choice not in ("1", "2", "3"):
            print("Error: Invalid data type selected.")
            return

        selected_type = getDataTypeName(type_choice)

        print("\nChoose the data source for the chart")
        print("1 Initial Data")
        print("2 Converted Data")

        data_choice = input("Choose an option: ")

        if data_choice == "1":
            data_column = 4

            if selected_type == "Temperature":
                y_axis_label = "Temperature (Fahrenheit)"

            elif selected_type == "Weight":
                y_axis_label = "Weight (Pounds)"

            else:
                y_axis_label = "Rain Amount (Inches)"

        elif data_choice == "2":
            data_column = 5

            if selected_type == "Temperature":
                y_axis_label = "Temperature (Celsius)"

            elif selected_type == "Weight":
                y_axis_label = "Weight (Kilograms)"

            else:
                y_axis_label = "Rain Amount (Centimeters)"

        else:
            print("Error: Invalid data source selected.")
            return

        rows = []

        with open(file_path, "r", newline="") as file:
            csv_reader = csv.reader(file)

            # Skip the header row.
            next(csv_reader, None)

            for row in csv_reader:
                if len(row) >= 6 and row[2] == selected_type:
                    rows.append([
                        row[0],
                        row[1],
                        row[3],
                        float(row[data_column])
                    ])

        if len(rows) == 0:
            print(
                "Error: No",
                selected_type,
                "data was found."
            )
            return

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        worksheet.append([
            "Date",
            "Location",
            "Description",
            y_axis_label
        ])

        for row in rows:
            worksheet.append(row)

        if chart_type == "bar":
            chart = BarChart()

        elif chart_type == "line":
            chart = LineChart()

        else:
            print("Error: Invalid chart type.")
            return

        chart.title = (
            "JadPow5144 "
            + datetime.now().strftime("%m/%d/%Y")
        )

        chart.x_axis.title = "Date"
        chart.y_axis.title = y_axis_label

        chart_data = Reference(
            worksheet,
            min_col=4,
            min_row=1,
            max_row=worksheet.max_row
        )

        chart_dates = Reference(
            worksheet,
            min_col=1,
            min_row=2,
            max_row=worksheet.max_row
        )

        chart.add_data(
            chart_data,
            titles_from_data=True
        )

        chart.set_categories(chart_dates)

        worksheet.add_chart(chart, "F2")

        workbook.save("final.xlsx")

        print("\nChart created successfully.")
        print("Data type:", selected_type)
        print("Measurement:", y_axis_label)
        print("The report was saved as final.xlsx.")

    except FileNotFoundError:
        print("Error: The CSV file could not be found.")

    except ValueError:
        print(
            "Error: The CSV file contains invalid numerical data."
        )

    except OSError as error:
        print("Error creating report:", error)


# generateReport requires file_path (string).
# It asks for a chart type, calls createChart, and returns no value.
def generateReport(file_path):
    print("\nChoose a graph type")
    print("1 Line Chart")
    print("2 Bar Chart")

    graph_choice = input("Choose an option: ")

    if graph_choice == "1":
        createChart(file_path, "line")

    elif graph_choice == "2":
        createChart(file_path, "bar")

    else:
        print("Error: Invalid graph type selected.")


# displayMenu requires no arguments.
# It displays the main menu and returns the user's selection.
def displayMenu():
    menu_options = [
        "Input Data",
        "View Current Data",
        "Generate Report",
        "Exit Program"
    ]

    print("\nJadPow5144's Spreadsheet Automation Menu")
    print("Choose a number from the following options")

    for number, option in enumerate(
        menu_options,
        start=1
    ):
        print(number, option)

    return input("Choose an option: ")


# main requires no arguments.
# It controls the program menu and returns no value.
def main():
    while True:

        selection = displayMenu()

        if selection == "1":
            print(
                "You selected",
                selection,
                "at",
                datetime.now()
            )

            getInput()

        elif selection == "2":
            print(
                "You selected",
                selection,
                "at",
                datetime.now()
            )

            viewData("ZooData.csv")

        elif selection == "3":
            print(
                "You selected",
                selection,
                "at",
                datetime.now()
            )

            generateReport("ZooData.csv")

        elif selection == "4":
            print(
                "You selected",
                selection,
                "at",
                datetime.now()
            )

            print("Program ended.")
            break

        else:
            print("Error: Invalid choice selected.")


main()
